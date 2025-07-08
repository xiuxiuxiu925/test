#!/usr/bin/env python3
"""
Test suite for Word Document Format Modifier
"""

import os
import sys
import tempfile
import shutil
from pathlib import Path
from unittest.mock import Mock, patch

# Add the parent directory to the path to import our module
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

try:
    from docx import Document
    from docx.shared import Pt
    from word_format_modifier import (
        FontManager, DocumentProcessor, ProgressTracker, 
        FontCategory, parse_color, find_word_files
    )
except ImportError as e:
    print(f"Error importing modules: {e}")
    print("Please install required packages: pip install -r requirements.txt")
    sys.exit(1)


class TestFontManager:
    """Test cases for FontManager"""
    
    def __init__(self):
        self.font_manager = FontManager()
    
    def test_font_initialization(self):
        """Test font database initialization"""
        assert len(self.font_manager.fonts) > 0
        assert "微软雅黑" in self.font_manager.fonts
        assert "Arial" in self.font_manager.fonts
        assert "等线" in self.font_manager.fonts
        print("✓ Font initialization test passed")
    
    def test_font_alias_mapping(self):
        """Test font alias mapping"""
        # Test Chinese font aliases
        assert self.font_manager.normalize_font_name("Microsoft YaHei") == "微软雅黑"
        assert self.font_manager.normalize_font_name("DengXian") == "等线"
        assert self.font_manager.normalize_font_name("SimSun") == "宋体"
        
        # Test English font aliases
        assert self.font_manager.normalize_font_name("Arial") == "Arial"
        assert self.font_manager.normalize_font_name("Times") == "Times New Roman"
        
        print("✓ Font alias mapping test passed")
    
    def test_font_suggestions(self):
        """Test font suggestions by category"""
        chinese_fonts = self.font_manager.get_font_suggestions(FontCategory.CHINESE)
        english_fonts = self.font_manager.get_font_suggestions(FontCategory.ENGLISH)
        
        assert len(chinese_fonts) > 0
        assert len(english_fonts) > 0
        assert "微软雅黑" in chinese_fonts
        assert "Arial" in english_fonts
        
        print("✓ Font suggestions test passed")
    
    def test_best_font_match(self):
        """Test best font matching with fallback"""
        # Test exact match
        assert self.font_manager.get_best_font_match("Arial") == "Arial"
        
        # Test fallback for non-existent font
        fallback = self.font_manager.get_best_font_match("NonExistentFont", FontCategory.CHINESE)
        assert fallback in self.font_manager.fonts
        
        print("✓ Best font match test passed")
    
    def run_all_tests(self):
        """Run all font manager tests"""
        print("Running FontManager tests...")
        self.test_font_initialization()
        self.test_font_alias_mapping()
        self.test_font_suggestions()
        self.test_best_font_match()
        print("All FontManager tests passed!\n")


class TestDocumentProcessor:
    """Test cases for DocumentProcessor"""
    
    def __init__(self):
        self.font_manager = FontManager()
        self.processor = DocumentProcessor(self.font_manager)
        self.temp_dir = None
    
    def setup_temp_directory(self):
        """Setup temporary directory for testing"""
        self.temp_dir = Path(tempfile.mkdtemp())
        return self.temp_dir
    
    def cleanup_temp_directory(self):
        """Clean up temporary directory"""
        if self.temp_dir and self.temp_dir.exists():
            shutil.rmtree(self.temp_dir)
    
    def create_test_document(self, file_path: Path, content: str = "Test document content"):
        """Create a test Word document"""
        doc = Document()
        doc.add_paragraph(content)
        doc.add_paragraph("Second paragraph with different content")
        doc.add_paragraph("Third paragraph for testing")
        doc.save(file_path)
        return file_path
    
    def test_backup_creation(self):
        """Test backup file creation"""
        temp_dir = self.setup_temp_directory()
        try:
            # Create test document
            test_file = temp_dir / "test.docx"
            self.create_test_document(test_file)
            
            # Create backup
            backup_path = self.processor.create_backup(test_file)
            
            assert backup_path.exists()
            assert backup_path.parent.name == "backups"
            assert "test_" in backup_path.name
            
            print("✓ Backup creation test passed")
            
        finally:
            self.cleanup_temp_directory()
    
    def test_file_lock_detection(self):
        """Test file lock detection"""
        temp_dir = self.setup_temp_directory()
        try:
            # Create test document
            test_file = temp_dir / "test.docx"
            self.create_test_document(test_file)
            
            # Test with unlocked file
            assert not self.processor.is_file_locked(test_file)
            
            print("✓ File lock detection test passed")
            
        finally:
            self.cleanup_temp_directory()
    
    def test_document_processing(self):
        """Test basic document processing"""
        temp_dir = self.setup_temp_directory()
        try:
            # Create test document
            test_file = temp_dir / "test.docx"
            self.create_test_document(test_file)
            
            # Process document
            modifications = {
                'font_name': 'Arial',
                'font_size': 14,
                'bold': True
            }
            
            result = self.processor.process_document(test_file, modifications)
            
            assert result['success']
            assert result['total_paragraphs'] > 0
            assert result['processing_time'] > 0
            
            print("✓ Document processing test passed")
            
        finally:
            self.cleanup_temp_directory()
    
    def test_batch_processing(self):
        """Test batch document processing"""
        temp_dir = self.setup_temp_directory()
        try:
            # Create multiple test documents
            test_files = []
            for i in range(3):
                test_file = temp_dir / f"test_{i}.docx"
                self.create_test_document(test_file, f"Test document {i}")
                test_files.append(test_file)
            
            # Process batch
            modifications = {'font_name': 'Arial', 'font_size': 12}
            results = self.processor.process_documents_batch(test_files, modifications, max_workers=2)
            
            assert len(results) == 3
            assert all(r['success'] for r in results)
            
            print("✓ Batch processing test passed")
            
        finally:
            self.cleanup_temp_directory()
    
    def run_all_tests(self):
        """Run all document processor tests"""
        print("Running DocumentProcessor tests...")
        self.test_backup_creation()
        self.test_file_lock_detection()
        self.test_document_processing()
        self.test_batch_processing()
        print("All DocumentProcessor tests passed!\n")


class TestUtilityFunctions:
    """Test utility functions"""
    
    def test_color_parsing(self):
        """Test color parsing function"""
        # Test color names
        assert parse_color("red") == (255, 0, 0)
        assert parse_color("green") == (0, 255, 0)
        assert parse_color("blue") == (0, 0, 255)
        
        # Test hex colors
        assert parse_color("#FF0000") == (255, 0, 0)
        assert parse_color("#00FF00") == (0, 255, 0)
        
        # Test RGB tuples
        assert parse_color("255,0,0") == (255, 0, 0)
        assert parse_color("0, 255, 0") == (0, 255, 0)
        
        # Test invalid colors
        assert parse_color("invalid") is None
        assert parse_color("#GGGGGG") is None
        
        print("✓ Color parsing test passed")
    
    def test_file_finding(self):
        """Test Word file finding function"""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            
            # Create test files
            (temp_path / "test1.docx").touch()
            (temp_path / "test2.doc").touch()
            (temp_path / "test3.txt").touch()  # Should be ignored
            (temp_path / "~$temp.docx").touch()  # Should be ignored (temp file)
            
            # Create subdirectory
            subdir = temp_path / "subdir"
            subdir.mkdir()
            (subdir / "test4.docx").touch()
            
            # Test non-recursive
            files = find_word_files(temp_path, recursive=False)
            assert len(files) == 2  # test1.docx and test2.doc
            
            # Test recursive
            files = find_word_files(temp_path, recursive=True)
            assert len(files) == 3  # Including test4.docx from subdir
            
            print("✓ File finding test passed")
    
    def run_all_tests(self):
        """Run all utility function tests"""
        print("Running utility function tests...")
        self.test_color_parsing()
        self.test_file_finding()
        print("All utility function tests passed!\n")


class TestProgressTracker:
    """Test progress tracking functionality"""
    
    def test_progress_tracker(self):
        """Test progress tracker creation and updates"""
        # Create mock progress tracker (to avoid actual progress bar output during tests)
        with patch('word_format_modifier.tqdm') as mock_tqdm:
            mock_progress_bar = Mock()
            mock_tqdm.return_value = mock_progress_bar
            
            tracker = ProgressTracker(100, "Test Progress")
            
            # Test update
            tracker.update(10)
            mock_progress_bar.update.assert_called_with(10)
            
            # Test description change
            tracker.set_description("New Description")
            mock_progress_bar.set_description.assert_called_with("New Description")
            
            # Test close
            tracker.close()
            mock_progress_bar.close.assert_called_once()
            
            print("✓ Progress tracker test passed")
    
    def run_all_tests(self):
        """Run all progress tracker tests"""
        print("Running ProgressTracker tests...")
        self.test_progress_tracker()
        print("All ProgressTracker tests passed!\n")


def run_all_tests():
    """Run all test suites"""
    print("=" * 60)
    print("Running Word Document Format Modifier Test Suite")
    print("=" * 60)
    
    try:
        # Run all test suites
        font_manager_tests = TestFontManager()
        font_manager_tests.run_all_tests()
        
        document_processor_tests = TestDocumentProcessor()
        document_processor_tests.run_all_tests()
        
        utility_tests = TestUtilityFunctions()
        utility_tests.run_all_tests()
        
        progress_tests = TestProgressTracker()
        progress_tests.run_all_tests()
        
        print("=" * 60)
        print("🎉 All tests passed successfully!")
        print("=" * 60)
        
    except Exception as e:
        print(f"❌ Test failed: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)


if __name__ == "__main__":
    run_all_tests()